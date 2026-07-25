# -*- coding: utf-8 -*-
"""
generate_test_functions_doc.py — Genera Funciones_de_Prueba_Optimizador.docx
y .pdf: qué función matemática se usó para probar cada método, para cuáles
métodos aplica y para cuáles no (incluye los casos probados deliberadamente
para verificar que la app bloquea combinaciones incompatibles).

Reutiliza el mismo estilo/setup que generate_docs.py (Tarea 4).

Uso:
    python scripts/generate_test_functions_doc.py
"""
from __future__ import annotations
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scripts.test_functions_data import (       # noqa: E402
    TEST_FUNCTIONS, INCOMPATIBLE_TESTS,
    INVENTORY_TESTS, INVENTORY_INCOMPATIBLE_TESTS,
)

OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
os.makedirs(OUT_DIR, exist_ok=True)
LOGO_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "Menu", "fondo_optimizacion.png")


def generate_docx(path: str) -> None:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    if os.path.exists(LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO_PATH, width=Cm(3.2))
    t = doc.add_heading("Optimizador de Funciones", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s = doc.add_heading("Funciones de Prueba por Método", level=1)
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"Generado automáticamente — {date.today().isoformat()}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Este documento lista las funciones matemáticas usadas para probar "
        "los 32 métodos del Optimizador durante la sesión de control de "
        "calidad, indicando con qué métodos cada función es compatible y "
        "con cuáles se verificó deliberadamente que la aplicación bloquea "
        "la ejecución.")
    doc.add_page_break()

    doc.add_heading("1. Funciones usadas para verificar ejecución correcta", level=1)
    for i, f in enumerate(TEST_FUNCTIONS, start=1):
        doc.add_heading(f"1.{i}  {f['funcion']}", level=2)
        for label, key in (("Parámetros", "parametros"), ("Resultado", "resultado")):
            para = doc.add_paragraph()
            r = para.add_run(f"{label}: ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            para.add_run(f[key])
        para = doc.add_paragraph()
        r = para.add_run("Aplica a: ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x2E, 0x8B, 0x57)
        para.add_run(", ".join(f["aplica_a"]))
        para = doc.add_paragraph()
        r = para.add_run("NO aplica a: ")
        r.bold = True
        r.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        para.add_run(" · ".join(f["no_aplica_a"]))

    doc.add_page_break()
    doc.add_heading("2. Funciones usadas para verificar el bloqueo de "
                     "combinaciones incompatibles", level=1)
    for i, t2 in enumerate(INCOMPATIBLE_TESTS, start=1):
        doc.add_heading(f"2.{i}  {t2['funcion']}  →  {t2['metodo']}", level=2)
        for label, key in (("Motivo de incompatibilidad", "motivo"),
                            ("Resultado en la app", "resultado_app")):
            para = doc.add_paragraph()
            r = para.add_run(f"{label}: ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            para.add_run(t2[key])

    doc.add_page_break()
    doc.add_heading("3. Módulo de Inventario — modelos usados para probar cada uno", level=1)
    for i, f in enumerate(INVENTORY_TESTS, start=1):
        doc.add_heading(f"3.{i}  {f['funcion']}", level=2)
        for label, key in (("Parámetros", "parametros"), ("Resultado", "resultado")):
            para = doc.add_paragraph()
            r = para.add_run(f"{label}: ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            para.add_run(f[key])
        para = doc.add_paragraph()
        r = para.add_run("Aplica a: ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x2E, 0x8B, 0x57)
        para.add_run(", ".join(f["aplica_a"]))
        para = doc.add_paragraph()
        r = para.add_run("NO aplica a: ")
        r.bold = True
        r.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        para.add_run(" · ".join(f["no_aplica_a"]))

    doc.add_page_break()
    doc.add_heading("4. Inventario — validaciones y bloqueos verificados", level=1)
    for i, t3 in enumerate(INVENTORY_INCOMPATIBLE_TESTS, start=1):
        doc.add_heading(f"4.{i}  {t3['funcion']}  →  {t3['metodo']}", level=2)
        for label, key in (("Motivo", "motivo"), ("Resultado en la app", "resultado_app")):
            para = doc.add_paragraph()
            r = para.add_run(f"{label}: ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            para.add_run(t3[key])

    doc.save(path)


def generate_pdf(path: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                     PageBreak, Image as RLImg)
    from xml.sax.saxutils import escape as _xesc

    styles = getSampleStyleSheet()
    h_title = ParagraphStyle("h_title", parent=styles["Title"], fontSize=20,
                              textColor=colors.HexColor("#1F4E79"))
    h1 = ParagraphStyle("h1c", parent=styles["Heading1"], fontSize=14,
                         textColor=colors.HexColor("#1F4E79"), spaceBefore=12, spaceAfter=6)
    h2 = ParagraphStyle("h2c", parent=styles["Heading2"], fontSize=11.5,
                         textColor=colors.HexColor("#2E75B6"), spaceBefore=9, spaceAfter=3)
    body = ParagraphStyle("bodyc", parent=styles["BodyText"], fontSize=9.3, leading=12.5)

    from scripts.pdf_cover import build_cover_elements

    M = 1.8 * cm
    doc = SimpleDocTemplate(path, pagesize=A4,
                             leftMargin=M, rightMargin=M,
                             topMargin=M, bottomMargin=M)
    elems = []
    elems += build_cover_elements(
        LOGO_PATH, "Optimizador de Funciones", "Funciones de Prueba por Método",
        intro_text=(
            "Este documento lista las funciones matemáticas usadas para probar "
            "los 32 métodos del Optimizador durante la sesión de control de "
            "calidad, indicando con qué métodos cada función es compatible y "
            "con cuáles se verificó deliberadamente que la aplicación bloquea "
            "la ejecución."),
        date_str=date.today().isoformat(), page_size=A4,
        top_margin=M, bottom_margin=M)
    elems.append(PageBreak())

    elems.append(Paragraph("1. Funciones usadas para verificar ejecución correcta", h1))
    for i, f in enumerate(TEST_FUNCTIONS, start=1):
        elems.append(Paragraph(f"1.{i}  {_xesc(f['funcion'])}", h2))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Parámetros:</font></b> {_xesc(f['parametros'])}", body))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Resultado:</font></b> {_xesc(f['resultado'])}", body))
        elems.append(Paragraph(f"<b><font color='#2E8B57'>Aplica a:</font></b> {_xesc(', '.join(f['aplica_a']))}", body))
        elems.append(Paragraph(f"<b><font color='#B03030'>NO aplica a:</font></b> {_xesc(' · '.join(f['no_aplica_a']))}", body))
        elems.append(Spacer(1, 0.25*cm))
    elems.append(PageBreak())

    elems.append(Paragraph("2. Funciones usadas para verificar el bloqueo de "
                            "combinaciones incompatibles", h1))
    for i, t2 in enumerate(INCOMPATIBLE_TESTS, start=1):
        elems.append(Paragraph(f"2.{i}  {_xesc(t2['funcion'])}  &rarr;  {_xesc(t2['metodo'])}", h2))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Motivo:</font></b> {_xesc(t2['motivo'])}", body))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Resultado en la app:</font></b> {_xesc(t2['resultado_app'])}", body))
        elems.append(Spacer(1, 0.25*cm))
    elems.append(PageBreak())

    elems.append(Paragraph("3. Módulo de Inventario — modelos usados para probar cada uno", h1))
    for i, f in enumerate(INVENTORY_TESTS, start=1):
        elems.append(Paragraph(f"3.{i}  {_xesc(f['funcion'])}", h2))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Parámetros:</font></b> {_xesc(f['parametros'])}", body))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Resultado:</font></b> {_xesc(f['resultado'])}", body))
        elems.append(Paragraph(f"<b><font color='#2E8B57'>Aplica a:</font></b> {_xesc(', '.join(f['aplica_a']))}", body))
        elems.append(Paragraph(f"<b><font color='#B03030'>NO aplica a:</font></b> {_xesc(' · '.join(f['no_aplica_a']))}", body))
        elems.append(Spacer(1, 0.25*cm))
    elems.append(PageBreak())

    elems.append(Paragraph("4. Inventario — validaciones y bloqueos verificados", h1))
    for i, t3 in enumerate(INVENTORY_INCOMPATIBLE_TESTS, start=1):
        elems.append(Paragraph(f"4.{i}  {_xesc(t3['funcion'])}  &rarr;  {_xesc(t3['metodo'])}", h2))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Motivo:</font></b> {_xesc(t3['motivo'])}", body))
        elems.append(Paragraph(f"<b><font color='#1F4E79'>Resultado en la app:</font></b> {_xesc(t3['resultado_app'])}", body))
        elems.append(Spacer(1, 0.25*cm))

    doc.build(elems)


def main() -> None:
    docx_path = os.path.join(OUT_DIR, "Funciones_de_Prueba_Optimizador.docx")
    pdf_path = os.path.join(OUT_DIR, "Funciones_de_Prueba_Optimizador.pdf")
    generate_docx(docx_path)
    print(f"OK docx: {docx_path} ({os.path.getsize(docx_path)} bytes)")
    generate_pdf(pdf_path)
    print(f"OK pdf:  {pdf_path} ({os.path.getsize(pdf_path)} bytes)")


if __name__ == "__main__":
    main()
